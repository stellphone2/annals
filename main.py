import calendar
import json
import re
import os
import shutil
import ssl
import subprocess
import sys
import base64
import time
import threading
import traceback
import urllib.error
import urllib.parse
import urllib.request
import webbrowser
from datetime import date, datetime, timedelta

def _global_exception_handler(exctype, value, tb):
    err_str = "".join(traceback.format_exception(exctype, value, tb))
    for path in ['/sdcard/Download/annals_crash.txt',
                 os.path.join(os.environ.get('ANDROID_PRIVATE', ''), 'error_log.txt'),
                 os.path.join(os.path.dirname(__file__), 'error_log.txt')]:
        try:
            with open(path, 'w', encoding='utf-8') as f:
                f.write(err_str)
        except Exception:
            pass
    sys.__excepthook__(exctype, value, tb)

sys.excepthook = _global_exception_handler

# Kivy Framework Imports
import kivy

kivy.require('2.0.0')

from kivy.app import App
from kivy.clock import Clock
from kivy.uix.boxlayout import BoxLayout
from kivy.uix.gridlayout import GridLayout
from kivy.uix.button import Button
from kivy.uix.dropdown import DropDown
from kivy.uix.spinner import Spinner
from kivy.uix.label import Label
from kivy.uix.textinput import TextInput
from kivy.uix.popup import Popup
from kivy.utils import platform

# Export Modules (loaded lazily on demand in export functions)

# Storage Locations
APP_DIR = os.path.dirname(os.path.abspath(__file__))

if platform == 'android':
    BASE_DIR = os.environ.get('ANDROID_PRIVATE', '')
    if not BASE_DIR or not os.path.exists(BASE_DIR):
        try:
            from android.storage import app_storage_path
            BASE_DIR = app_storage_path()
        except Exception:
            BASE_DIR = APP_DIR
else:
    BASE_DIR = APP_DIR

NOTES_FILE = os.path.join(BASE_DIR, 'annals_notes.json')
EXCEL_FILE = os.path.join(BASE_DIR, 'annals_notes.xlsx')
DOCX_FILE = os.path.join(BASE_DIR, 'annals_notes.docx')
PDF_FILE = os.path.join(BASE_DIR, 'annals_notes.pdf')
ERROR_LOG_FILE = os.path.join(BASE_DIR, 'error_log.txt')
CLOUD_SETTINGS_FILE = os.path.join(BASE_DIR, 'cloud_settings.json')  # stays on the device only

# The ONLY cloud setting kept in the app: the Apps Script web app URL (ends in /exec).
# Everything else (Google client, who may do what) lives in the Google script / Access sheet.
CLOUD_URL = 'https://script.google.com/macros/s/AKfycbweS4YcUGXYR0Ww3fg-tGhrdw58gi3p9l-8RCRMbPcyTJHuUhQfF4Iu4UBo-7W7Papq/exec'
ACCESS_REFRESH_SECONDS = 300  # how often the community list / notes are re-read from the sheet
NO_COMMUNITIES_TEXT = 'Sign in to load communities'
REFRESH_LIST_TEXT = 'Refresh list'
ALL_COMMUNITIES_LABEL = 'All communities'
STORE_VERSION = 2


# Core Helper Functions
def log_error(error_message, exc_info=None):
    timestamp = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    try:
        with open(ERROR_LOG_FILE, 'a', encoding='utf-8') as f:
            f.write(f'[{timestamp}] {error_message}\n')
            if exc_info:
                traceback.print_exception(*exc_info, file=f)
            f.write('-' * 60 + '\n')
    except Exception as e:
        print(f'Failed to log error to file: {e}')


def parse_note_data(val):
    """Extracts (text, community) from the legacy v1 note format."""
    if val is None:
        return '', 'PLT'
    if isinstance(val, dict):
        return str(val.get('text', '')).strip(), str(val.get('community', 'PLT')).strip() or 'PLT'
    return str(val).strip(), 'PLT'


# File Opening Helper
def open_file_with_default_app(file_path):
    """Opens/shares a generated file. Returns True if a viewer/share sheet was launched."""
    if not os.path.exists(file_path):
        return False
    try:
        if platform == 'android':
            # Needs `androidstorage4kivy` in buildozer requirements. It handles
            # FileProvider URIs, so we avoid the FileUriExposedException that
            # Uri.fromFile() causes on Android 7+.
            from androidstorage4kivy import SharedStorage, ShareSheet
            try:
                # Keep a copy in shared storage so the user can find it later.
                SharedStorage().copy_to_shared(file_path)
            except Exception:
                log_error('Copy to shared storage failed', sys.exc_info())
            ShareSheet().share_file(file_path)
            return True
        elif platform == 'win':
            os.startfile(os.path.abspath(file_path))
        elif platform == 'macosx':
            subprocess.call(['open', file_path])
        else:
            webbrowser.open('file://' + os.path.abspath(file_path))
        return True
    except Exception:
        log_error(f'Failed to open generated file: {file_path}', sys.exc_info())
        return False


# File-lock helper
def _try_release_file(file_path):
    """
    Windows-only: if *file_path* is locked by another process, attempt to
    close the known viewer / editor that holds it, then verify the file is
    writable.  Returns True when the file can be written (or does not yet
    exist).  Always returns True on non-Windows platforms.
    """
    if not os.path.exists(file_path):
        return True          # nothing to release
    if platform != 'win':
        return True          # locking only matters on Windows

    def _writable():
        try:
            with open(file_path, 'r+b'):
                return True
        except PermissionError:
            return False

    if _writable():
        return True          # file is not locked

    # Map each extension to the process names of apps that lock it
    ext = os.path.splitext(file_path)[1].lower()
    viewer_procs = {
        '.pdf':  ['AcroRd32.exe', 'Acrobat.exe', 'SumatraPDF.exe',
                  'FoxitReader.exe', 'FoxitPDFReader.exe'],
        '.xlsx': ['EXCEL.EXE'],
        '.docx': ['WINWORD.EXE'],
    }
    for proc_name in viewer_procs.get(ext, []):
        try:
            subprocess.run(
                ['taskkill', '/f', '/im', proc_name],
                capture_output=True, timeout=5
            )
        except Exception:
            pass

    time.sleep(0.5)          # give Windows time to release handles
    return _writable()


# Cloud (Google Sheet through an Apps Script web app)
CLOUD_URL_PREFIX = 'https://script.google.com/'
MAX_CELL_CHARS = 49000  # Google Sheets allows 50,000 characters per cell
SETTINGS_KEYS = ('password', 'mode', 'client_id', 'client_secret', 'refresh_token', 'email')   # url is the constant above


def _read_settings_file():
    try:
        with open(CLOUD_SETTINGS_FILE, 'r', encoding='utf-8') as f:
            data = json.load(f)
        return data if isinstance(data, dict) else {}
    except FileNotFoundError:
        return {}
    except Exception:
        log_error('Error loading cloud settings', sys.exc_info())
        return {}


def load_cloud_settings():
    """dict with url, client_id, client_secret (both fetched from the script), refresh_token, email."""
    data = _read_settings_file()
    s = {k: str(data.get(k, '')).strip() for k in SETTINGS_KEYS}
    s['url'] = CLOUD_URL.strip()
    s['password_ok'] = s['password'] == today_password()
    # 'email' = simple mode (the person types their Gmail); anything else = real Google sign-in
    s['signed_in'] = bool(s['email'] if s['mode'] == 'email' else s['refresh_token'])
    return s


def save_cloud_settings(**changes):
    """Merge the given keys into the settings file (stays on the device only)."""
    data = _read_settings_file()
    data.update({k: v for k, v in changes.items() if k in SETTINGS_KEYS})
    try:
        with open(CLOUD_SETTINGS_FILE, 'w', encoding='utf-8') as f:
            json.dump(data, f)
        return True
    except Exception:
        log_error('Error saving cloud settings', sys.exc_info())
        return False


def _ssl_context():
    try:
        import certifi  # Python on Android has no usable system CA store
        return ssl.create_default_context(cafile=certifi.where())
    except Exception:
        return ssl.create_default_context()


PASSWORD_PREFIX = '20'    # must be the same in annals_cloud.gs
PASSWORD_NEEDED = 'password-needed'


def today_password():
    """Daily cloud password: PASSWORD_PREFIX + today's DDMMYY (India time), e.g. 5 Oct 2026 -> 20051026."""
    ist = datetime.utcnow() + timedelta(hours=5, minutes=30)
    return PASSWORD_PREFIX + ist.strftime('%d%m%y')


def _post_json(url, payload, timeout=45):
    """POST JSON to the web app. Returns (reply_dict, None) or (None, error_message)."""
    payload = dict(payload, password=str(_read_settings_file().get('password', '')))
    data = json.dumps(payload, ensure_ascii=False).encode('utf-8')
    req = urllib.request.Request(
        url, data=data, method='POST',
        headers={'Content-Type': 'application/json; charset=utf-8'})
    try:
        with urllib.request.urlopen(req, timeout=timeout, context=_ssl_context()) as resp:
            body = resp.read().decode('utf-8', 'replace')
    except urllib.error.HTTPError as e:
        return None, f'The web app returned HTTP {e.code}.'
    except Exception as e:  # offline, DNS, timeout, SSL ...
        return None, f'Could not reach Google:\n{e}'
    try:
        reply = json.loads(body)
    except ValueError:
        reply = None
    if not isinstance(reply, dict):
        return None, ('Unexpected reply from the web app.\n'
                      'Check that it is deployed with "Who has access: Anyone"\n'
                      'and that the URL ends in /exec.')
    return reply, None


# --- Google sign-in (OAuth "device" flow: the person enters a short code on google.com/device) ---
GOOGLE_DEVICE_URL = 'https://oauth2.googleapis.com/device/code'
GOOGLE_TOKEN_URL = 'https://oauth2.googleapis.com/token'
SIGN_IN_NEEDED = 'sign-in-needed'
_id_cache = {'token': '', 'exp': 0, 'rt': ''}
_token_lock = threading.Lock()


def _post_form(url, fields, timeout=30):
    """POST form fields to Google. Google reports problems as 4xx with a JSON body, so read that too."""
    req = urllib.request.Request(
        url, data=urllib.parse.urlencode(fields).encode('utf-8'), method='POST',
        headers={'Content-Type': 'application/x-www-form-urlencoded'})
    try:
        with urllib.request.urlopen(req, timeout=timeout, context=_ssl_context()) as resp:
            body = resp.read().decode('utf-8', 'replace')
    except urllib.error.HTTPError as e:
        body = e.read().decode('utf-8', 'replace')
    except Exception as e:
        return None, f'Could not reach Google:\n{e}'
    try:
        reply = json.loads(body)
    except ValueError:
        reply = None
    if not isinstance(reply, dict):
        return None, 'Unexpected reply from Google.'
    return reply, None


def _google_problem(reply):
    return str(reply.get('error_description') or reply.get('error') or 'Unknown error from Google.')


def id_token_info(token):
    """Claims of a Google ID token (email, exp ...). The web app does the real verification."""
    try:
        part = token.split('.')[1]
        part += '=' * (-len(part) % 4)
        info = json.loads(base64.urlsafe_b64decode(part.encode('ascii')).decode('utf-8'))
        return info if isinstance(info, dict) else {}
    except Exception:
        return {}


def _remember_tokens(reply, refresh_token):
    token = str(reply.get('id_token') or '')
    if token:
        _id_cache.update(token=token, exp=float(id_token_info(token).get('exp') or 0), rt=refresh_token)


def google_device_start(client_id):
    """Returns (info, None) with user_code / verification_url / device_code, or (None, message)."""
    reply, err = _post_form(GOOGLE_DEVICE_URL, {'client_id': client_id, 'scope': 'openid email'})
    if err:
        return None, err
    if reply.get('error') or not reply.get('device_code'):
        return None, _google_problem(reply)
    reply.setdefault('verification_url', reply.get('verification_uri', 'https://www.google.com/device'))
    return reply, None


def google_device_poll(client_id, client_secret, device_code):
    """One poll. Returns ('pending'|'slow'|'done'|'error', reply_or_message)."""
    reply, err = _post_form(GOOGLE_TOKEN_URL, {
        'client_id': client_id, 'client_secret': client_secret, 'device_code': device_code,
        'grant_type': 'urn:ietf:params:oauth:grant-type:device_code'})
    if err:
        return 'pending', err           # a network hiccup: keep trying until the code expires
    problem = reply.get('error')
    if problem == 'authorization_pending':
        return 'pending', ''
    if problem == 'slow_down':
        return 'slow', ''
    if problem:
        return 'error', _google_problem(reply)
    if not reply.get('id_token') or not reply.get('refresh_token'):
        return 'error', 'Google did not return a sign-in. Check the OAuth client type and scopes.'
    return 'done', reply


def get_id_token():
    """(id_token, None) or (None, error). Error SIGN_IN_NEEDED means the person must sign in again."""
    s = load_cloud_settings()
    if s['mode'] == 'email':
        return (f'email:{s["email"]}', None) if s['email'] else (None, SIGN_IN_NEEDED)
    if not (s['client_id'] and s['client_secret'] and s['refresh_token']):
        return None, SIGN_IN_NEEDED
    with _token_lock:
        if (_id_cache['token'] and _id_cache['rt'] == s['refresh_token']
                and _id_cache['exp'] - time.time() > 120):
            return _id_cache['token'], None
        reply, err = _post_form(GOOGLE_TOKEN_URL, {
            'client_id': s['client_id'], 'client_secret': s['client_secret'],
            'refresh_token': s['refresh_token'], 'grant_type': 'refresh_token'})
        if err:
            return None, err
        if reply.get('error') in ('invalid_grant', 'invalid_client', 'unauthorized_client'):
            return None, SIGN_IN_NEEDED
        if reply.get('error') or not reply.get('id_token'):
            return None, _google_problem(reply)
        _remember_tokens(reply, s['refresh_token'])
        return reply['id_token'], None


def sign_out():
    _id_cache.update(token='', exp=0, rt='')
    return save_cloud_settings(refresh_token='', email='')


def open_url(url):
    """Open a web page in the phone's / computer's browser."""
    try:
        if platform == 'android':
            from jnius import autoclass
            Intent = autoclass('android.content.Intent')
            Uri = autoclass('android.net.Uri')
            activity = autoclass('org.kivy.android.PythonActivity').mActivity
            activity.startActivity(Intent(Intent.ACTION_VIEW, Uri.parse(url)))
        else:
            webbrowser.open(url)
    except Exception:
        log_error('Could not open the browser', sys.exc_info())


OUTDATED_SCRIPT_MSG = ('Your Google Apps Script is out of date.\n'
                       'Paste the latest annals_cloud.gs and deploy a New version.')


def _reply_error(reply):
    msg = str(reply.get('error', 'Unknown error from the web app.'))
    if reply.get('auth') is False and 'password' in msg.lower():
        return PASSWORD_NEEDED
    if reply.get('auth') is False and ('expired' in msg or 'Not signed in' in msg or 'requires Google' in msg):
        return SIGN_IN_NEEDED
    if 'not listed' in msg:
        return msg + '\nCheck the address with Export Notes > Cloud sign out.'
    if 'Unknown action' in msg or 'Wrong or missing token' in msg:
        return OUTDATED_SCRIPT_MSG
    return msg


def clean_email(text):
    """Lower-cased Gmail address, or '' if it does not look like one."""
    t = str(text).strip().lower()
    return t if re.fullmatch(r'[^@\s,;]+@[^@\s,;]+\.[^@\s,;]+', t) else ''


def fetch_config(url):
    """How the script wants people to identify themselves.
    Returns ({'mode': 'email'|'google', 'client_id', 'client_secret'}, None) or (None, error).
    A change of mode or Google client signs everybody out once."""
    reply, err = _post_json(url, {'action': 'config'})
    if err:
        return None, err
    if not reply.get('ok'):
        return None, str(reply.get('error') or OUTDATED_SCRIPT_MSG)
    mode = 'email' if reply.get('mode') == 'email' else 'google'
    cid, secret = str(reply.get('client_id', '')).strip(), str(reply.get('client_secret', '')).strip()
    if mode == 'google' and not (cid and secret):
        return None, OUTDATED_SCRIPT_MSG
    old = _read_settings_file()
    if (mode, cid, secret) != (old.get('mode'), old.get('client_id', ''), old.get('client_secret', '')):
        save_cloud_settings(mode=mode, client_id=cid, client_secret=secret, refresh_token='', email='')
        _id_cache.update(token='', exp=0, rt='')
    return {'mode': mode, 'client_id': cid, 'client_secret': secret}, None


def fetch_access(url, id_token):
    """The Access worksheet. Returns ([(name, can_write), ...], email, None) or (None, '', error)."""
    reply, err = _post_json(url, {'id_token': id_token, 'action': 'access'})
    if err:
        return None, '', err
    if not reply.get('ok'):
        return None, '', _reply_error(reply)
    raw = reply.get('communities')
    if not isinstance(raw, list):
        return None, '', OUTDATED_SCRIPT_MSG
    out = []
    for c in raw:
        if isinstance(c, dict) and str(c.get('name', '')).strip():
            out.append((str(c['name']).strip(), bool(c.get('write'))))
    return out, str(reply.get('email', '')), None


def parse_sheet_rows(raw_rows):
    """Sheet rows [[date, note], ...] -> ({'YYYY-MM-DD': note}, skipped_rows)."""
    notes, skipped = {}, 0
    for row in raw_rows:
        if not isinstance(row, (list, tuple)) or len(row) < 2:
            skipped += 1
            continue
        date_str = str(row[0]).strip()
        note = str(row[1]).replace('\r\n', '\n').replace('\r', '\n').strip()
        if not date_str and not note:
            continue                                   # blank line
        if date_str.lower() == 'date' and note.lower() == 'note':
            continue                                   # header row
        try:
            date_str = datetime.strptime(date_str, '%Y-%m-%d').date().isoformat()  # '2026-1-7' -> '2026-01-07'
        except ValueError:
            skipped += 1
            continue
        if not note:
            continue
        # two rows for the same date: keep both texts
        notes[date_str] = f'{notes[date_str]}\n\n{note}' if date_str in notes else note
    return notes, skipped


def save_to_cloud(url, id_token, sheet_name, rows, timeout=45):
    """Save [(date, full_text), ...]: new dates are appended as a row, an existing date's row is
    extended (the script refuses anything that does not keep the old text as its beginning).
    Returns (ok, message, saved_dates). Never raises."""
    reply, err = _post_json(url, {
        'id_token': id_token,
        'action': 'save',
        'sheet': sheet_name,
        'rows': [[d, n[:MAX_CELL_CHARS]] for d, n in rows],
    }, timeout)
    if err:
        return False, err, []
    if not reply.get('ok'):
        return False, _reply_error(reply), []
    saved, skipped = reply.get('saved'), reply.get('skipped') or []
    if not isinstance(saved, list):
        return False, OUTDATED_SCRIPT_MSG, []
    msg = f'Saved {len(saved)} notes to\nworksheet "{sheet_name}".'
    if skipped:
        msg += f'\n{len(skipped)} could not be saved:'
        for item in skipped[:5]:
            msg += f'\n{item.get("date")}: {item.get("reason")}'
    return True, msg, [str(d) for d in saved]


class SheetNotes(dict):
    """{date: note} from the sheet, plus what the script said about the worksheet."""
    row_count = 0
    tab_found = None
    tabs = ()


def pull_from_cloud(url, id_token, community, timeout=45):
    """Read one community's worksheet. Returns (notes_dict, skipped_rows, error_or_None).
    The request carries no 'sheet' key, so an out-of-date script rejects it
    instead of treating it as an upload."""
    reply, err = _post_json(url, {'id_token': id_token, 'action': 'read', 'community': community}, timeout)
    if err:
        return None, 0, err
    if not reply.get('ok'):
        return None, 0, _reply_error(reply)
    if not isinstance(reply.get('rows'), list):
        return None, 0, OUTDATED_SCRIPT_MSG
    parsed, skipped = parse_sheet_rows(reply['rows'])
    notes = SheetNotes(parsed)
    notes.row_count = len(reply['rows'])
    notes.tab_found = reply.get('found')               # None = script too old to say
    notes.tabs = [str(t) for t in reply.get('sheets') or []]
    return notes, skipped, None


def describe_sync(stats, skipped=0, sheet_notes=None, community=''):
    parts = [f'{stats[k]} {k}' for k in ('added', 'updated') if stats[k]]
    line = 'From the sheet: ' + (', '.join(parts) if parts else 'no changes') + '.'
    if stats['merged']:
        line += (f'\n{stats["merged"]} of your notes already had a row in the sheet;'
                 '\nyour text is added after it.')
    if skipped:
        line += f'\n{skipped} sheet row(s) skipped (bad date).'
    if sheet_notes is not None and getattr(sheet_notes, 'tab_found', None) is False:
        tabs = ', '.join(f'"{t}"' for t in sheet_notes.tabs) or 'none'
        line += (f'\nNo worksheet named "{community}" exists in the sheet (names must match exactly).'
                 f'\nWorksheets found: {tabs}.')
    elif sheet_notes is not None and not sheet_notes:
        line += f'\nThe worksheet "{community}" has no usable notes.'
    elif sheet_notes is not None:
        line += f'\nSheet has {len(sheet_notes)} dated notes.'
    return line


# Data Store
class NoteStore:
    """
    On-disk format (v2):
    {
      "version": 2,
      "communities": ["PLT", "Dum Dum"],   # copy of the sheet's Access list (works offline)
      "writable": ["PLT"],                 # the ones this person may add notes to
      "selected": "PLT",
      "notes": {"2026-10-03": {"PLT": "text", "Dum Dum": "text"}}
    }
    Legacy v1 files (date -> text or {text, community}) are migrated automatically
    and the original is kept as <file>.v1.bak.
    """

    def __init__(self, path):
        self.path = path
        self.communities = []          # copy of the Access sheet; never edited in the app
        self.writable = set()          # communities this person may add notes to
        self.selected = ''
        self.notes = {}
        self.uploaded = {}   # community -> {date: text already saved to the cloud (read-only)}
        self._load()

    # --- loading / saving ---
    def _load(self):
        if not os.path.exists(self.path):
            return
        try:
            with open(self.path, 'r', encoding='utf-8') as f:
                data = json.load(f)
        except Exception:
            log_error('Error loading notes file', sys.exc_info())
            self._backup(self.path + '.corrupt-' + datetime.now().strftime('%Y%m%d-%H%M%S'))
            return

        if isinstance(data, dict) and data.get('version') == STORE_VERSION:
            self._load_v2(data)
        else:
            self._backup(self.path + '.v1.bak')
            self._migrate_v1(data)
            self.save()

    def _backup(self, dest):
        if not os.path.exists(self.path):
            return
        try:
            shutil.copy2(self.path, dest)
        except Exception:
            log_error('Could not back up notes file', sys.exc_info())

    def _load_v2(self, data):
        comms = []
        for c in data.get('communities', []):
            c = str(c).strip()
            if c and c not in comms:
                comms.append(c)

        notes = {}
        raw_notes = data.get('notes', {})
        if isinstance(raw_notes, dict):
            for date_key, per_comm in raw_notes.items():
                if not isinstance(per_comm, dict):
                    continue
                for comm, text in per_comm.items():
                    comm, text = str(comm).strip(), str(text).strip()
                    if comm and text:
                        notes.setdefault(date_key, {})[comm] = text

        self.notes = notes
        self.communities = comms
        raw_writable = data.get('writable', [])
        self.writable = {str(c) for c in raw_writable if str(c) in comms} if isinstance(raw_writable, list) else set()
        selected = str(data.get('selected', '')).strip()
        self.selected = selected if selected in comms else (comms[0] if comms else '')

        uploaded = {}
        raw_uploaded = data.get('uploaded', {})
        if isinstance(raw_uploaded, dict):
            for comm, per in raw_uploaded.items():
                if isinstance(per, dict):
                    uploaded[str(comm)] = {str(d): str(t) for d, t in per.items()}
                elif isinstance(per, list):   # earlier format: just the dates; the note text is the saved text
                    uploaded[str(comm)] = {str(d): self.notes[str(d)][str(comm)] for d in per
                                           if str(comm) in self.notes.get(str(d), {})}
        self.uploaded = uploaded

    def _migrate_v1(self, data):
        comms = []
        notes = {}
        if isinstance(data, dict):
            for date_key, val in data.items():
                text, comm = parse_note_data(val)
                if text:
                    notes.setdefault(date_key, {})[comm] = text
                    if comm not in comms:
                        comms.append(comm)
        self.notes = notes
        self.communities = comms          # replaced by the sheet's list at the first sign-in
        self.selected = comms[0] if comms else ''

    def save(self):
        tmp_path = self.path + '.tmp'
        try:
            payload = {
                'version': STORE_VERSION,
                'communities': self.communities,
                'writable': sorted(self.writable),
                'selected': self.selected,
                'notes': self.notes,
                'uploaded': self.uploaded,
            }
            with open(tmp_path, 'w', encoding='utf-8') as f:
                json.dump(payload, f, ensure_ascii=False, indent=2)
                f.flush()
                os.fsync(f.fileno())
            os.replace(tmp_path, self.path)  # atomic: never leaves a half-written file
            return True
        except Exception:
            log_error('Error saving notes file', sys.exc_info())
            return False

    # --- notes ---
    def get(self, d, community):
        return self.notes.get(d.isoformat(), {}).get(community, '')

    def has(self, d, community):
        return bool(self.get(d, community))

    def can_write(self, community):
        return community in self.writable

    def set(self, d, community, text):
        if not self.can_write(community):
            return True          # view-only community: nothing can be added here
        text = text.strip()
        saved = self.uploaded_text(d, community)
        if saved is not None:
            # saved text is read-only: only a note that still begins with it is accepted
            if not text.startswith(saved):
                return True
            if text == self.get(d, community):
                return True
            self.notes.setdefault(d.isoformat(), {})[community] = text
            return self.save()
        if not text:
            return self.delete(d, community)
        self.notes.setdefault(d.isoformat(), {})[community] = text
        return self.save()

    def delete(self, d, community):
        if not self.can_write(community) or self.is_locked(d, community):
            return True          # saved to the cloud: cannot be deleted
        key = d.isoformat()
        per_comm = self.notes.get(key)
        if per_comm and community in per_comm:
            del per_comm[community]
            if not per_comm:
                del self.notes[key]
            return self.save()
        return True

    def query(self, start_d=None, end_d=None, community=None):
        rows = []
        for date_str, per_comm in self.notes.items():
            try:
                d_obj = datetime.strptime(date_str, '%Y-%m-%d').date()
            except ValueError:
                log_error(f'Bad date key in notes: {date_str}')
                continue
            if start_d and d_obj < start_d:
                continue
            if end_d and d_obj > end_d:
                continue
            for comm, text in per_comm.items():
                if community and comm != community:
                    continue
                if text:
                    rows.append((date_str, comm, text))
        rows.sort(key=lambda r: (r[0], r[1].lower()))
        return rows

    # --- cloud (what is saved can only be added to) ---
    def uploaded_text(self, d, community):
        """The text of this day that is already in the cloud (read-only), or None."""
        return self.uploaded.get(community, {}).get(d.isoformat())

    def is_locked(self, d, community):
        return self.uploaded_text(d, community) is not None

    def is_pending(self, d, community):
        """True when the note has something that is not in the cloud yet."""
        text = self.get(d, community)
        return bool(text) and text != self.uploaded_text(d, community)

    def has_pending(self, community):
        """True when any note of this community has text that is not in the cloud yet."""
        saved = self.uploaded.get(community, {})
        return any(text != saved.get(d) for d, _c, text in self.query(community=community))

    def has_uploaded(self, community):
        return bool(self.uploaded.get(community))

    def pending_rows(self, community, sheet_notes):
        """[(date, full_text), ...] of this community that differ from what the sheet holds."""
        return [(d, text) for d, _c, text in self.query(community=community)
                if sheet_notes.get(d) != text]

    def sync_from_sheet(self, community, sheet_notes):
        """Bring the sheet's text into the app. Whatever the sheet holds becomes the read-only
        part of that day; anything typed here that is not in the sheet stays after it.
          date only in the sheet            -> added
          same text on both sides           -> just marked as uploaded (e.g. a reply that got lost)
          sheet changed since our upload    -> read-only part replaced, our additions kept after it
          not uploaded yet, sheet has a row -> sheet text first, ours added after it ('merged')
        Returns counts: added, updated, merged."""
        prefixes = self.uploaded.setdefault(community, {})
        before = dict(prefixes)
        stats = {'added': 0, 'updated': 0, 'merged': 0}
        new_text = {}
        for d, S in sheet_notes.items():
            L, P = self.notes.get(d, {}).get(community), prefixes.get(d)
            if L is None:
                new_text[d] = S
                prefixes[d] = S
                stats['added'] += 1
            elif L == S:
                prefixes[d] = S
            elif P is None:
                new_text[d] = f'{S}\n\n{L}'
                prefixes[d] = S
                stats['merged'] += 1
            elif S != P:
                extra = (L[len(P):].strip() if L.startswith(P) else L)
                new_text[d] = f'{S}\n\n{extra}' if extra else S
                prefixes[d] = S
                stats['updated'] += 1
            # else: the sheet is as we left it and we hold an addition -> nothing to do
        if new_text:
            self._backup(self.path + '.pre-cloud.bak')
            for d, text in new_text.items():
                self.notes.setdefault(d, {})[community] = text
        if new_text or prefixes != before:
            self.save()
        return stats

    def mark_uploaded(self, community, saved_rows):
        """saved_rows: [(date, text), ...] that the sheet now holds."""
        prefixes = self.uploaded.setdefault(community, {})
        for d, text in saved_rows:
            prefixes[d] = text[:MAX_CELL_CHARS]
        return self.save()

    # --- communities (the list comes from the Access sheet; the app cannot add or rename) ---
    def set_communities(self, names, writable):
        """Replace the list with the sheet's. Notes of communities that left the list stay in the
        file but are not shown. Returns True when anything changed."""
        clean = []
        for n in names:
            n = str(n).strip()
            if n and n not in clean:
                clean.append(n)
        can = {n for n in writable if n in clean}
        changed = clean != self.communities or can != self.writable
        self.communities, self.writable = clean, can
        if self.selected not in clean:
            self.selected = clean[0] if clean else ''
            changed = True
        if changed:
            self.save()
        return changed

    def set_selected(self, name):
        if name in self.communities and name != self.selected:
            self.selected = name
            self.save()


# Word + PDF builders (pure Python: no python-docx, lxml, pillow or fpdf needed)
_XML_BAD_CHARS = re.compile('[\x00-\x08\x0b\x0c\x0e-\x1f]')


def _xml_escape(text):
    text = _XML_BAD_CHARS.sub('', str(text))
    return text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')


def _docx_paragraph(text, bold=False, italic=False, size=None, after=0):
    """One <w:p>. Newlines in text become line breaks inside the same paragraph."""
    props = ''
    if bold:
        props += '<w:b/>'
    if italic:
        props += '<w:i/>'
    if size:
        props += f'<w:sz w:val="{size * 2}"/><w:szCs w:val="{size * 2}"/>'
    rpr = f'<w:rPr>{props}</w:rPr>' if props else ''
    runs = []
    for i, line in enumerate(str(text).replace('\r\n', '\n').replace('\r', '\n').split('\n')):
        br = '<w:br/>' if i else ''
        runs.append(f'<w:r>{rpr}{br}<w:t xml:space="preserve">{_xml_escape(line)}</w:t></w:r>')
    ppr = f'<w:pPr><w:spacing w:after="{after}"/></w:pPr>'
    return f'<w:p>{ppr}{"".join(runs)}</w:p>'


def _docx_cell(text, width, bold=False):
    return (f'<w:tc><w:tcPr><w:tcW w:w="{width}" w:type="dxa"/></w:tcPr>'
            f'{_docx_paragraph(text, bold=bold)}</w:tc>')


def build_docx(path, start_d, end_d, community, rows):
    """Write a .docx (Title, date-range line, Date/[Community]/Note table) using only zipfile."""
    import zipfile
    single = community is not None
    widths = [1800, 7946] if single else [1700, 1900, 6146]      # twips; A4 with 0.75" margins
    labels = ('Date', 'Note') if single else ('Date', 'Community', 'Note')

    border = ''.join(f'<w:{side} w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
                     for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'))
    table = [f'<w:tbl><w:tblPr><w:tblW w:w="{sum(widths)}" w:type="dxa"/>'
             f'<w:tblBorders>{border}</w:tblBorders><w:tblLayout w:type="fixed"/>'
             '<w:tblCellMar><w:left w:w="80" w:type="dxa"/><w:right w:w="80" w:type="dxa"/></w:tblCellMar>'
             '</w:tblPr><w:tblGrid>' + ''.join(f'<w:gridCol w:w="{w}"/>' for w in widths) + '</w:tblGrid>']
    table.append('<w:tr><w:trPr><w:tblHeader/></w:trPr>' +
                 ''.join(_docx_cell(l, w, bold=True) for l, w in zip(labels, widths)) + '</w:tr>')
    for date_str, comm, note in rows:
        cells = [date_str, note] if single else [date_str, comm, note]
        table.append('<w:tr>' + ''.join(_docx_cell(c, w) for c, w in zip(cells, widths)) + '</w:tr>')
    table.append('</w:tbl>')

    body = (_docx_paragraph('Annals', bold=True, size=18, after=60) +
            _docx_paragraph(f'Date Range: {start_d} to {end_d}  |  {community or ALL_COMMUNITIES_LABEL}',
                            italic=True, size=10, after=200) +
            ''.join(table) + _docx_paragraph(''))
    document = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                f'<w:body>{body}<w:sectPr><w:pgSz w:w="11906" w:h="16838"/>'
                '<w:pgMar w:top="1080" w:right="1080" w:bottom="1080" w:left="1080" '
                'w:header="708" w:footer="708" w:gutter="0"/></w:sectPr></w:body></w:document>')
    content_types = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                     '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                     '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                     '<Default Extension="xml" ContentType="application/xml"/>'
                     '<Override PartName="/word/document.xml" '
                     'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
                     '</Types>')
    rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
            'Target="word/document.xml"/></Relationships>')
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as z:
        z.writestr('[Content_Types].xml', content_types)
        z.writestr('_rels/.rels', rels)
        z.writestr('word/document.xml', document)


# Minimal PDF writer (pure Python, English text): the same code runs on a computer and on Android.
# Standard Helvetica needs no font files; these are its character widths (1/1000 em) for codes 32..255.
_HELV_R = [
    278, 278, 355, 556, 556, 889, 667, 191, 333, 333, 389, 584, 278, 333, 278, 278,
    556, 556, 556, 556, 556, 556, 556, 556, 556, 556, 278, 278, 584, 584, 584, 556,
    1015, 667, 667, 722, 722, 667, 611, 778, 722, 278, 500, 667, 556, 833, 722, 778,
    667, 778, 722, 667, 611, 722, 667, 944, 667, 667, 611, 278, 278, 278, 469, 556,
    333, 556, 556, 500, 556, 556, 278, 556, 556, 222, 222, 500, 222, 833, 556, 556,
    556, 556, 333, 500, 278, 556, 500, 722, 500, 500, 500, 334, 260, 334, 584, 350,
    556, 350, 222, 556, 333, 1000, 556, 556, 333, 1000, 667, 333, 1000, 350, 611, 350,
    350, 222, 222, 333, 333, 350, 556, 1000, 333, 1000, 500, 333, 944, 350, 500, 667,
    278, 333, 556, 556, 556, 556, 260, 556, 333, 737, 370, 556, 584, 333, 737, 333,
    400, 584, 333, 333, 333, 556, 537, 278, 333, 333, 365, 556, 834, 834, 834, 611,
    667, 667, 667, 667, 667, 667, 1000, 722, 667, 667, 667, 667, 278, 278, 278, 278,
    722, 722, 778, 778, 778, 778, 778, 584, 778, 722, 722, 722, 722, 667, 667, 611,
    556, 556, 556, 556, 556, 556, 889, 500, 556, 556, 556, 556, 278, 278, 278, 278,
    556, 556, 556, 556, 556, 556, 556, 584, 611, 556, 556, 556, 556, 500, 556, 500,
]

_HELV_B = [
    278, 333, 474, 556, 556, 889, 722, 238, 333, 333, 389, 584, 278, 333, 278, 278,
    556, 556, 556, 556, 556, 556, 556, 556, 556, 556, 333, 333, 584, 584, 584, 611,
    975, 722, 722, 722, 722, 667, 611, 778, 722, 278, 556, 722, 611, 833, 722, 778,
    667, 778, 722, 667, 611, 722, 667, 944, 667, 667, 611, 333, 278, 333, 584, 556,
    333, 556, 611, 556, 611, 556, 333, 611, 611, 278, 278, 556, 278, 889, 611, 611,
    611, 611, 389, 556, 333, 611, 556, 778, 556, 556, 500, 389, 280, 389, 584, 350,
    556, 350, 278, 556, 500, 1000, 556, 556, 333, 1000, 667, 333, 1000, 350, 611, 350,
    350, 278, 278, 500, 500, 350, 556, 1000, 333, 1000, 556, 333, 944, 350, 500, 667,
    278, 333, 556, 556, 556, 556, 280, 556, 333, 737, 370, 556, 584, 333, 737, 333,
    400, 584, 333, 333, 333, 611, 556, 278, 333, 333, 365, 556, 834, 834, 834, 611,
    722, 722, 722, 722, 722, 722, 1000, 722, 667, 667, 667, 667, 278, 278, 278, 278,
    722, 722, 778, 778, 778, 778, 778, 584, 778, 722, 722, 722, 722, 667, 667, 611,
    556, 556, 556, 556, 556, 556, 889, 556, 556, 556, 556, 556, 278, 278, 278, 278,
    611, 611, 611, 611, 611, 611, 611, 584, 611, 611, 611, 611, 611, 556, 611, 556,
]

_HELV_I = [
    278, 278, 355, 556, 556, 889, 667, 191, 333, 333, 389, 584, 278, 333, 278, 278,
    556, 556, 556, 556, 556, 556, 556, 556, 556, 556, 278, 278, 584, 584, 584, 556,
    1015, 667, 667, 722, 722, 667, 611, 778, 722, 278, 500, 667, 556, 833, 722, 778,
    667, 778, 722, 667, 611, 722, 667, 944, 667, 667, 611, 278, 278, 278, 469, 556,
    333, 556, 556, 500, 556, 556, 278, 556, 556, 222, 222, 500, 222, 833, 556, 556,
    556, 556, 333, 500, 278, 556, 500, 722, 500, 500, 500, 334, 260, 334, 584, 350,
    556, 350, 222, 556, 333, 1000, 556, 556, 333, 1000, 667, 333, 1000, 350, 611, 350,
    350, 222, 222, 333, 333, 350, 556, 1000, 333, 1000, 500, 333, 944, 350, 500, 667,
    278, 333, 556, 556, 556, 556, 260, 556, 333, 737, 370, 556, 584, 333, 737, 333,
    400, 584, 333, 333, 333, 556, 537, 278, 333, 333, 365, 556, 834, 834, 834, 611,
    667, 667, 667, 667, 667, 667, 1000, 722, 667, 667, 667, 667, 278, 278, 278, 278,
    722, 722, 778, 778, 778, 778, 778, 584, 778, 722, 722, 722, 722, 667, 667, 611,
    556, 556, 556, 556, 556, 556, 889, 500, 556, 556, 556, 556, 278, 278, 278, 278,
    556, 556, 556, 556, 556, 556, 556, 584, 611, 556, 556, 556, 556, 500, 556, 500,
]


_PDF_PAGE_W, _PDF_PAGE_H = 595.28, 841.89      # A4 in points
_PDF_MARGIN = 42.5
_PDF_FONT_SIZE = 10
_PDF_LINE_H = 12.5
_PDF_PAD = 4


def _pdf_clean(text):
    """Text safe for the standard PDF fonts: Latin-1 only, anything else becomes '?'."""
    text = str(text).replace('\r\n', '\n').replace('\r', '\n').replace('\t', '    ')
    swaps = {'‘': "'", '’': "'", '“': '"', '”': '"',
             '–': '-', '—': '-', '…': '...', '•': '*'}
    out = []
    for ch in text:
        o = ord(ch)
        if ch == '\n' or 32 <= o <= 126 or 160 <= o <= 255:
            out.append(ch)
        else:
            out.append(swaps.get(ch, '?'))
    return ''.join(out)


def _pdf_width(text, table, size):
    return sum(table[ord(c) - 32] for c in text) * size / 1000.0


def _pdf_wrap(text, table, size, max_w):
    """Word-wrap text into lines that fit max_w points (very long words are split)."""
    lines = []
    for para in _pdf_clean(text).split('\n'):
        cur = ''
        for word in para.split(' '):
            cand = word if not cur else cur + ' ' + word
            if _pdf_width(cand, table, size) <= max_w:
                cur = cand
                continue
            if cur:
                lines.append(cur)
                cur = ''
            while len(word) > 1 and _pdf_width(word, table, size) > max_w:
                k = 1
                while k < len(word) and _pdf_width(word[:k + 1], table, size) <= max_w:
                    k += 1
                lines.append(word[:k])
                word = word[k:]
            cur = word
        lines.append(cur)
    return lines or ['']


def _pdf_esc(text):
    return text.replace('\\', '\\\\').replace('(', '\\(').replace(')', '\\)')


class _PdfPages:
    """Collects drawing commands page by page; write() produces the PDF file."""

    def __init__(self):
        self.pages = []
        self.ops = None

    def new_page(self):
        self.ops = []
        self.pages.append(self.ops)

    def text(self, x, y, text, font='F1', size=_PDF_FONT_SIZE, gray=0.0, rgb=None):
        color = f'{rgb[0]:.3f} {rgb[1]:.3f} {rgb[2]:.3f} rg' if rgb else f'{gray:.3f} g'
        self.ops.append(f'BT {color} /{font} {size} Tf {x:.2f} {y:.2f} Td ({_pdf_esc(text)}) Tj ET')

    def rect(self, x, y, w, h):
        self.ops.append(f'0.5 w 0 G {x:.2f} {y:.2f} {w:.2f} {h:.2f} re S')

    def write(self, path):
        import zlib
        total = len(self.pages)
        for n, ops in enumerate(self.pages, 1):                     # footer: Page n/N
            label = f'Page {n}/{total}'
            w = _pdf_width(label, _HELV_R, 8)
            ops.append(f'BT 0.590 g /F3 8 Tf {(_PDF_PAGE_W - w) / 2:.2f} 24 Td ({label}) Tj ET')
        objs = {
            1: b'<< /Type /Catalog /Pages 2 0 R >>',
            3: b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>',
            4: b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold /Encoding /WinAnsiEncoding >>',
            5: b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Oblique /Encoding /WinAnsiEncoding >>',
        }
        kids, next_id = [], 6
        for ops in self.pages:
            page_id, content_id = next_id, next_id + 1
            next_id += 2
            data = zlib.compress('\n'.join(ops).encode('latin-1', 'replace'))
            objs[content_id] = (b'<< /Length %d /Filter /FlateDecode >>\nstream\n' % len(data)
                                + data + b'\nendstream')
            objs[page_id] = (f'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {_PDF_PAGE_W} {_PDF_PAGE_H}] '
                             f'/Resources << /Font << /F1 3 0 R /F2 4 0 R /F3 5 0 R >> >> '
                             f'/Contents {content_id} 0 R >>').encode('ascii')
            kids.append(f'{page_id} 0 R')
        objs[2] = f'<< /Type /Pages /Kids [{" ".join(kids)}] /Count {len(kids)} >>'.encode('ascii')

        out = bytearray(b'%PDF-1.4\n%\xe2\xe3\xcf\xd3\n')
        offsets = {}
        for num in sorted(objs):
            offsets[num] = len(out)
            out += b'%d 0 obj\n' % num + objs[num] + b'\nendobj\n'
        xref = len(out)
        out += b'xref\n0 %d\n0000000000 65535 f \n' % (next_id)
        for num in range(1, next_id):
            out += b'%010d 00000 n \n' % offsets[num]
        out += (b'trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n' % (next_id, xref))
        tmp = path + '.tmp'
        with open(tmp, 'wb') as f:
            f.write(out)
        os.replace(tmp, path)


def build_pdf(path, start_str, end_str, community, rows):
    """Write the notes as a real PDF: title, date-range line and a Date/[Community]/Note table."""
    single = community is not None
    content_w = _PDF_PAGE_W - 2 * _PDF_MARGIN
    widths = [85, content_w - 85] if single else [70, 95, content_w - 165]
    labels = ('Date', 'Note') if single else ('Date', 'Community', 'Note')
    size, lh, pad = _PDF_FONT_SIZE, _PDF_LINE_H, _PDF_PAD
    bottom = _PDF_MARGIN + 12
    doc = _PdfPages()
    state = {'y': 0.0}

    def start_page():
        doc.new_page()
        top = _PDF_PAGE_H - _PDF_MARGIN
        doc.text(_PDF_MARGIN, top - 14, 'Annals', font='F2', size=16, rgb=(0.173, 0.243, 0.314))
        subtitle = _pdf_clean(f'Date Range: {start_str} to {end_str}  |  {community or ALL_COMMUNITIES_LABEL}')
        doc.text(_PDF_MARGIN, top - 30, subtitle, font='F3', size=10, rgb=(0.498, 0.549, 0.553))
        y = top - 42
        x = _PDF_MARGIN
        for w, label in zip(widths, labels):                         # table header row
            doc.rect(x, y - 18, w, 18)
            doc.text(x + pad, y - 13, label, font='F2')
            x += w
        state['y'] = y - 18

    def draw_row(cols, height):
        y_top = state['y']
        x = _PDF_MARGIN
        for w, lines in zip(widths, cols):
            doc.rect(x, y_top - height, w, height)
            for i, line in enumerate(lines):
                if line:
                    doc.text(x + pad, y_top - pad - 8 - i * lh, line)
            x += w
        state['y'] = y_top - height

    start_page()
    for date_str, comm, note in rows:
        date_lines = _pdf_wrap(date_str, _HELV_R, size, widths[0] - 2 * pad)
        note_lines = _pdf_wrap(note, _HELV_R, size, widths[-1] - 2 * pad)
        short = [date_lines]
        if not single:
            short.append(_pdf_wrap(comm, _HELV_R, size, widths[1] - 2 * pad))
        short_n = max(len(c) for c in short)

        first, idx, fresh = True, 0, False
        while True:
            room = int((state['y'] - bottom - 2 * pad) // lh)
            need = min(len(note_lines) - idx, 3)
            if first:
                need = max(need, short_n)
            if fresh:
                need = 1                    # never loop forever on an empty page
            if room < need:
                start_page()
                fresh = True
                continue
            seg = note_lines[idx:idx + room]
            n_lines = max(len(seg), short_n if first else 0)
            blank = [['']] * len(short)
            draw_row((short if first else blank) + [seg], n_lines * lh + 2 * pad)
            idx += len(seg)
            first, fresh = False, False
            if idx >= len(note_lines):
                break
    doc.write(path)


# Dialog & Message UI Popups
class AlertPopup(Popup):
    def __init__(self, title, text, **kwargs):
        super().__init__(**kwargs)
        self.title = title
        self.size_hint = (0.8, 0.4)

        layout = BoxLayout(orientation='vertical', padding=10, spacing=10)
        layout.add_widget(Label(text=text, halign='center', valign='middle'))

        ok_btn = Button(text='OK', size_hint_y=0.3)
        ok_btn.bind(on_release=self.dismiss)
        layout.add_widget(ok_btn)

        self.content = layout


class ConfirmPopup(Popup):
    def __init__(self, title, text, on_yes, **kwargs):
        super().__init__(**kwargs)
        self.title = title
        self.on_yes = on_yes
        self.size_hint = (0.8, 0.4)

        layout = BoxLayout(orientation='vertical', padding=10, spacing=10)
        layout.add_widget(Label(text=text, halign='center', valign='middle'))

        btn_box = BoxLayout(size_hint_y=0.3, spacing=10)
        no_btn = Button(text='Cancel')
        no_btn.bind(on_release=self.dismiss)
        yes_btn = Button(text='Delete')
        yes_btn.bind(on_release=self._confirm)
        btn_box.add_widget(no_btn)
        btn_box.add_widget(yes_btn)
        layout.add_widget(btn_box)

        self.content = layout

    def _confirm(self, instance):
        self.dismiss()
        self.on_yes()


class PromptPopup(Popup):
    """Asks for one line of text (Gmail address, daily password). check(text) returns the cleaned value or ''."""

    def __init__(self, title, message, hint, check, on_ok, on_cancel, secret=False, **kwargs):
        super().__init__(**kwargs)
        self.title = title
        self.auto_dismiss = False
        self.size_hint = (0.92, 0.5)
        self.check, self.on_ok, self.on_cancel = check, on_ok, on_cancel
        layout = BoxLayout(orientation='vertical', padding=10, spacing=10)
        layout.add_widget(Label(text=message, halign='center', size_hint_y=0.35))
        self.input = TextInput(hint_text=hint, multiline=False, password=secret, size_hint_y=0.3)
        layout.add_widget(self.input)
        btn_box = BoxLayout(size_hint_y=0.35, spacing=10)
        cancel_btn = Button(text='Cancel')
        cancel_btn.bind(on_release=self._cancel)
        ok_btn = Button(text='OK')
        ok_btn.bind(on_release=self._ok)
        btn_box.add_widget(cancel_btn)
        btn_box.add_widget(ok_btn)
        layout.add_widget(btn_box)
        self.content = layout

    def _cancel(self, instance):
        self.dismiss()
        self.on_cancel()

    def _ok(self, instance):
        value = self.check(self.input.text)
        if not value:
            self.title = 'Not correct - try again'
            return
        self.dismiss()
        self.on_ok(value)


class CredentialsPopup(Popup):
    """First-run screen: the daily password and the Gmail address, together on one screen."""

    def __init__(self, on_ok, on_cancel, **kwargs):
        super().__init__(**kwargs)
        self.title = 'Welcome to Annals'
        self.auto_dismiss = False
        self.size_hint = (0.92, 0.7)
        self.on_ok, self.on_cancel = on_ok, on_cancel

        layout = BoxLayout(orientation='vertical', padding=10, spacing=8)
        layout.add_widget(Label(text="Enter cloud password", size_hint_y=0.12))
        self.password_input = TextInput(hint_text='password', multiline=False, password=True,
                                        size_hint_y=0.17)
        layout.add_widget(self.password_input)
        self.email_label = Label(text='Enter valid email', size_hint_y=0.12)
        layout.add_widget(self.email_label)
        self.email_input = TextInput(hint_text='name@gmail.com', multiline=False,
                                     input_type='mail', write_tab=False, size_hint_y=0.17)
        layout.add_widget(self.email_input)
        self.error_label = Label(text='', color=(1, 0.4, 0.4, 1), size_hint_y=0.1)
        layout.add_widget(self.error_label)

        btn_box = BoxLayout(size_hint_y=0.2, spacing=10)
        cancel_btn = Button(text='Cancel')
        cancel_btn.bind(on_release=self._cancel)
        ok_btn = Button(text='OK')
        ok_btn.bind(on_release=self._ok)
        btn_box.add_widget(cancel_btn)
        btn_box.add_widget(ok_btn)
        layout.add_widget(btn_box)
        self.content = layout

    def _cancel(self, instance):
        self.dismiss()
        self.on_cancel()

    def _ok(self, instance):
        password = self.password_input.text.strip()
        email = clean_email(self.email_input.text)
        problems = []
        if password != today_password():
            problems.append('Password is not correct')
        if not email:
            problems.append('Enter valid email')
        if problems:
            self.error_label.text = ' | '.join(problems)
            return
        self.dismiss()
        self.on_ok(password, email)


class SignInPopup(Popup):
    """Shows the short code the person types at google.com/device."""

    def __init__(self, user_code, verification_url, **kwargs):
        super().__init__(**kwargs)
        self.title = 'Sign in with Google'
        self.cancelled = False
        self.verification_url = verification_url
        self.auto_dismiss = False
        self.size_hint = (0.92, 0.65)

        layout = BoxLayout(orientation='vertical', padding=10, spacing=10)
        layout.add_widget(Label(
            text=f'1. Tap "Open Google page"\n   (or go to {verification_url})\n'
                 '2. Enter this code and choose your Gmail account:',
            halign='center', valign='middle', size_hint_y=0.4))
        layout.add_widget(Label(text=user_code, font_size='34sp', bold=True, size_hint_y=0.25))
        layout.add_widget(Label(text='This window closes by itself when you are done.',
                                halign='center', size_hint_y=0.15))

        btn_box = BoxLayout(size_hint_y=0.2, spacing=10)
        cancel_btn = Button(text='Cancel')
        cancel_btn.bind(on_release=self._cancel)
        open_btn = Button(text='Open Google page')
        open_btn.bind(on_release=lambda x: open_url(self.verification_url))
        btn_box.add_widget(cancel_btn)
        btn_box.add_widget(open_btn)
        layout.add_widget(btn_box)

        self.content = layout

    def _cancel(self, instance):
        self.cancelled = True
        self.dismiss()


class DatePickerRow(BoxLayout):
    """Day / Month / Year spinners, so nobody has to type YYYY-MM-DD on a phone."""

    MONTHS = list(calendar.month_abbr)[1:]

    def __init__(self, initial, **kwargs):
        super().__init__(spacing=4, **kwargs)
        years = [str(y) for y in range(date.today().year - 10, date.today().year + 3)]
        if str(initial.year) not in years:
            years.append(str(initial.year))
            years.sort()

        self.day = Spinner(text=str(initial.day), values=[str(i) for i in range(1, 32)],
                           size_hint_x=0.25)
        self.month = Spinner(text=self.MONTHS[initial.month - 1], values=self.MONTHS,
                             size_hint_x=0.35)
        self.year = Spinner(text=str(initial.year), values=years, size_hint_x=0.4)
        for w in (self.day, self.month, self.year):
            self.add_widget(w)

    def get_date(self):
        year = int(self.year.text)
        month = self.MONTHS.index(self.month.text) + 1
        day = min(int(self.day.text), calendar.monthrange(year, month)[1])  # clamp 31 Feb etc.
        return date(year, month, day)


class DateRangePopup(Popup):
    def __init__(self, title, communities, default_community, callback, **kwargs):
        super().__init__(**kwargs)
        self.title = title
        self.callback = callback
        self.size_hint = (0.92, 0.6)

        layout = BoxLayout(orientation='vertical', padding=10, spacing=10)

        grid = GridLayout(cols=2, row_default_height=44, row_force_default=True, spacing=5)
        grid.add_widget(Label(text='Start Date:', size_hint_x=0.25))
        self.start_picker = DatePickerRow(date.today().replace(day=1))
        grid.add_widget(self.start_picker)

        grid.add_widget(Label(text='End Date:', size_hint_x=0.25))
        self.end_picker = DatePickerRow(date.today())
        grid.add_widget(self.end_picker)

        grid.add_widget(Label(text='Community:', size_hint_x=0.25))
        self.comm_spinner = Spinner(
            text=default_community if default_community in communities else ALL_COMMUNITIES_LABEL,
            values=[ALL_COMMUNITIES_LABEL] + list(communities),
        )
        grid.add_widget(self.comm_spinner)
        layout.add_widget(grid)

        btn_box = BoxLayout(size_hint_y=0.3, spacing=10)
        cancel_btn = Button(text='Cancel')
        cancel_btn.bind(on_release=self.dismiss)
        confirm_btn = Button(text='Export')
        confirm_btn.bind(on_release=self.validate_and_submit)

        btn_box.add_widget(cancel_btn)
        btn_box.add_widget(confirm_btn)
        layout.add_widget(btn_box)

        self.content = layout

    def validate_and_submit(self, instance):
        start_d = self.start_picker.get_date()
        end_d = self.end_picker.get_date()
        if start_d > end_d:
            AlertPopup('Date Error', 'Start Date cannot be after End Date.').open()
            return
        community = None if self.comm_spinner.text == ALL_COMMUNITIES_LABEL else self.comm_spinner.text
        self.dismiss()
        self.callback(start_d, end_d, community)


class NoteEditorPopup(Popup):
    """note_text is the part you can edit. uploaded_text is the part already saved to the
    cloud: it is shown read-only and new text can only be added after it."""

    def __init__(self, current_date, note_text, app_instance, uploaded_text=None, read_only=False,
                 **kwargs):
        super().__init__(**kwargs)
        self.read_only = read_only             # a community you may only view
        self.uploaded_text = uploaded_text
        self.locked = uploaded_text is not None or read_only
        self.current_date = current_date
        self.app_instance = app_instance
        self.title = f'Notes - {current_date.strftime("%d-%m-%Y")}'
        self.size_hint = (0.95, 0.9)
        self.auto_dismiss = False  # only Close/Prev/Next/Delete leave the editor, so nothing is lost

        layout = BoxLayout(orientation='vertical', padding=10, spacing=10)

        # Top Control Bar
        top_bar = BoxLayout(size_hint_y=0.1, spacing=5)
        btn_del = Button(text='Delete')
        btn_del.bind(on_release=self.confirm_delete)
        btn_del.disabled = self.locked          # saved text can never be deleted
        btn_prev = Button(text='< Prev')
        btn_prev.bind(on_release=lambda x: self.change_date(-1))
        btn_next = Button(text='Next >')
        btn_next.bind(on_release=lambda x: self.change_date(1))
        btn_close = Button(text='Close')
        btn_close.bind(on_release=self.close_and_save)

        top_bar.add_widget(btn_del)
        top_bar.add_widget(btn_prev)
        top_bar.add_widget(btn_next)
        layout.add_widget(top_bar)

        # Date & Community Header Display
        header = f"{current_date.strftime('%a, %d %b %Y')} ({self.app_instance.selected_community})"
        if read_only:
            header += '\nView only - you cannot add notes to this community'
        elif self.locked:
            header += '\nSaved to cloud - you can only add more below'
        layout.add_widget(
            Label(text=header, font_size='18sp', bold=True,
                  size_hint_y=0.12 if self.locked else 0.08, halign='center')
        )

        if read_only:
            self.text_input = TextInput(
                text=uploaded_text or '', readonly=True, multiline=True, font_size='16sp',
                size_hint_y=0.68, background_color=(0.88, 0.91, 1, 1))
            layout.add_widget(self.text_input)
        else:
            if self.locked:
                layout.add_widget(TextInput(
                    text=uploaded_text, readonly=True, multiline=True, font_size='16sp',
                    size_hint_y=0.3, background_color=(0.88, 0.91, 1, 1)))

            # Multi-line Text Area (the whole note, or only the addition when part of it is saved)
            self.text_input = TextInput(
                text=note_text, multiline=True, font_size='16sp',
                hint_text='Add more to this day...' if self.locked else '',
                size_hint_y=0.38 if self.locked else 0.72
            )
            layout.add_widget(self.text_input)

        # Bottom bar: Close
        layout.add_widget(btn_close)
        btn_close.size_hint_y = 0.1

        self.content = layout

    def _full_text(self):
        if not self.locked:
            return self.text_input.text
        extra = self.text_input.text.strip()
        return f'{self.uploaded_text}\n\n{extra}' if extra else self.uploaded_text

    def close_and_save(self, *args):
        if not self.read_only:
            self.app_instance.save_note(self.current_date, self._full_text())
        self.dismiss()

    def confirm_delete(self, instance):
        if self.locked:
            return
        if not self.text_input.text.strip():
            self.dismiss()
            return
        ConfirmPopup(
            'Delete Note',
            f'Delete this note for {self.current_date.strftime("%d-%m-%Y")}\n'
            f'({self.app_instance.selected_community})?',
            on_yes=self._do_delete,
        ).open()

    def _do_delete(self):
        self.app_instance.delete_note(self.current_date)
        self.dismiss()

    def change_date(self, direction):
        if not self.read_only:
            self.app_instance.save_note(self.current_date, self._full_text())
        self.dismiss()
        new_date = self.current_date + timedelta(days=direction)
        self.app_instance.open_notes(new_date)


# Main App Root Layout
class AnnalsNotesLayout(BoxLayout):
    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.orientation = 'vertical'
        self.padding = 5
        self.spacing = 5

        self.store = NoteStore(NOTES_FILE)
        self.current_year = date.today().year
        self.current_month = date.today().month
        self._last_today = date.today()
        self._updating_spinner = False
        self._editor_open = 0          # while a note is open, background updates wait
        self._signing_in = False
        self._asking_password = False
        self._sync_text = ''
        self._sync_error = False       # True while the status line shows a problem
        self._last_sync = ''           # 'HH:MM' of the last successful sync / refresh
        self._syncing = False          # a Sync is running: ignore extra taps and stale background reads
        self._sync_started = 0.0
        self._sync_gen = 0             # bumps at every Sync start/end; older background reads are dropped

        # Month Title Header
        self.month_label = Label(
            text='', font_size='20sp', bold=True, size_hint_y=0.08
        )
        self.add_widget(self.month_label)

        # Navigation Bar
        nav_bar = BoxLayout(size_hint_y=0.08, spacing=3)
        btn_prev = Button(text='< Prev')
        btn_prev.bind(on_release=lambda x: self.previous_month())
        btn_today = Button(text='Today')
        btn_today.bind(on_release=lambda x: self.go_today())
        btn_next = Button(text='Next >')
        btn_next.bind(on_release=lambda x: self.next_month())

        # Export Dropdown
        self.export_dropdown = DropDown()
        for label, kind in (('Export Excel', 'excel'), ('Export Word', 'docx'), ('Export PDF', 'pdf'),
                            ('Cloud sign out', 'signout')):
            b = Button(text=label, size_hint_y=None, height=44)
            b.bind(on_release=lambda btn, k=kind: self._handle_export_select(k))
            self.export_dropdown.add_widget(b)

        self.export_btn = Button(text='Export Notes')
        self.export_btn.bind(on_release=self.export_dropdown.open)

        nav_bar.add_widget(btn_prev)
        nav_bar.add_widget(btn_today)
        nav_bar.add_widget(btn_next)
        nav_bar.add_widget(self.export_btn)
        self.add_widget(nav_bar)

        # Community Dropdown Selector
        comm_bar = BoxLayout(size_hint_y=0.07, spacing=5)
        comm_bar.add_widget(Label(text='Community:', size_hint_x=0.28, bold=True))

        self.comm_spinner = Spinner(
            text=self.selected_community or NO_COMMUNITIES_TEXT,
            values=self._get_spinner_options(),
            size_hint_x=0.5
        )
        self.comm_spinner.bind(text=self.on_community_change)
        comm_bar.add_widget(self.comm_spinner)
        self.sync_btn = Button(text='Sync', size_hint_x=0.22)
        self.sync_btn.bind(on_release=lambda x: self._send_to_cloud())
        comm_bar.add_widget(self.sync_btn)
        self.add_widget(comm_bar)

        # Permission / sync status line
        self.status_label = Label(text='', font_size='13sp', size_hint_y=0.05)
        self.add_widget(self.status_label)

        # Calendar View Container
        self.calendar_grid = GridLayout(cols=7, spacing=2, size_hint_y=0.72)
        self.add_widget(self.calendar_grid)

        self.create_calendar()
        self._update_status()

        # First run: ask for password and email together on one screen
        self._asking_credentials = False
        Clock.schedule_once(self._first_run_check, 0.3)

        # The community list and the notes come from the Google Sheet
        Clock.schedule_once(lambda dt: self.refresh_access(silent=True), 1)
        Clock.schedule_interval(lambda dt: self.refresh_access(silent=True), ACCESS_REFRESH_SECONDS)

        # Refresh the "today" highlight if the app stays open past midnight
        Clock.schedule_interval(self._check_day_rollover, 60)

    # --- helpers ---
    @property
    def selected_community(self):
        return self.store.selected

    def _persist_failed_alert(self):
        AlertPopup('Save Error', 'Could not save your changes.\nSee error_log.txt for details.').open()

    def _get_spinner_options(self):
        return list(self.store.communities) + [REFRESH_LIST_TEXT]

    def _set_spinner(self, text):
        """Update the spinner without triggering on_community_change."""
        self._updating_spinner = True
        try:
            self.comm_spinner.values = self._get_spinner_options()
            self.comm_spinner.text = text or NO_COMMUNITIES_TEXT
        finally:
            self._updating_spinner = False

    def _check_day_rollover(self, dt):
        if date.today() != self._last_today:
            self._last_today = date.today()
            self.create_calendar()

    def _set_sync(self, text, error=False):
        self._sync_text = text
        self._sync_error = error
        self._update_status()

    def _mark_synced(self):
        self._last_sync = f'{datetime.now():%H:%M}'

    def _update_sync_button(self):
        """Red while something is not synced yet, green when everything is up to date."""
        community = self.selected_community
        pending = bool(community) and self.store.can_write(community) and self.store.has_pending(community)
        self.sync_btn.background_color = (0.85, 0.2, 0.2, 1) if pending else (0.2, 0.7, 0.3, 1)

    def _update_status(self):
        community = self.selected_community
        if not community:
            parts = ['No communities yet - use Sync to sign in']
        else:
            parts = []
            if not self.store.can_write(community):
                parts.append('View only')
            if self._sync_text:
                parts.append(self._sync_text)
            if self._last_sync:
                parts.append(f'Last sync {self._last_sync}')
        self.status_label.text = '  |  '.join(parts)

    # --- communities (the list comes from the Access sheet) ---
    def on_community_change(self, spinner, text):
        if self._updating_spinner:
            return
        if text == REFRESH_LIST_TEXT:
            self._set_spinner(self.selected_community)
            self.refresh_access(silent=False)
        elif text == NO_COMMUNITIES_TEXT:
            self._set_spinner(self.selected_community)
        else:
            self.store.set_selected(text)
            self.create_calendar()
            self._update_status()
            self._background_pull(text)

    # --- calendar ---
    def create_calendar(self):
        self.calendar_grid.clear_widgets()
        self.month_label.text = f'{calendar.month_name[self.current_month]} {self.current_year}'

        for day in ['Mon', 'Tue', 'Wed', 'Thu', 'Fri', 'Sat', 'Sun']:
            self.calendar_grid.add_widget(
                Label(text=day, bold=True, size_hint_y=None, height=30)
            )

        today = date.today()
        community = self.selected_community
        for week in calendar.monthcalendar(self.current_year, self.current_month):
            for day_num in week:
                if not day_num:
                    self.calendar_grid.add_widget(Label(text=''))
                    continue

                d = date(self.current_year, self.current_month, day_num)
                has_note = self.store.has(d, community)

                btn_text = f'{day_num}\n\u2022' if has_note else str(day_num)
                btn = Button(text=btn_text, font_size='14sp')

                if d == today:
                    btn.background_color = (0.2, 0.6, 1, 1)
                elif has_note and not self.store.is_pending(d, community):
                    btn.background_color = (0.55, 0.62, 0.85, 1)   # fully saved to the cloud
                elif has_note:
                    btn.background_color = (0.4, 0.8, 0.5, 1)

                btn.bind(on_release=lambda instance, dt=d: self.open_notes(dt))
                self.calendar_grid.add_widget(btn)
        self._update_sync_button()

    def previous_month(self):
        if self.current_month == 1:
            self.current_month = 12
            self.current_year -= 1
        else:
            self.current_month -= 1
        self.create_calendar()

    def next_month(self):
        if self.current_month == 12:
            self.current_month = 1
            self.current_year += 1
        else:
            self.current_month += 1
        self.create_calendar()

    def go_today(self):
        today = date.today()
        self.current_year = today.year
        self.current_month = today.month
        self.create_calendar()

    # --- notes ---
    def open_notes(self, d):
        community = self.selected_community
        if not community:
            AlertPopup('Annals', 'No communities yet.\nUse Sync and sign in\n'
                                 'to load the list from the Google Sheet.').open()
            return
        full = self.store.get(d, community)
        if not self.store.can_write(community):
            popup = NoteEditorPopup(current_date=d, note_text='', app_instance=self,
                                    uploaded_text=full, read_only=True)
        else:
            saved = self.store.uploaded_text(d, community)
            editable = full if saved is None else (full[len(saved):].strip() if full.startswith(saved) else '')
            popup = NoteEditorPopup(current_date=d, note_text=editable, app_instance=self,
                                    uploaded_text=saved)
        self._editor_open += 1
        popup.bind(on_dismiss=lambda *a: self._editor_closed())
        popup.open()
        if not popup.read_only:
            Clock.schedule_once(lambda dt: setattr(popup.text_input, 'focus', True), 0.2)

    def _editor_closed(self):
        self._editor_open = max(0, self._editor_open - 1)

    def save_note(self, d, text_content):
        if not self.store.set(d, self.selected_community, text_content):
            self._persist_failed_alert()
        self.create_calendar()

    def delete_note(self, d):
        if not self.store.delete(d, self.selected_community):
            self._persist_failed_alert()
        self.create_calendar()

    # --- exports ---
    def _deliver(self, path):
        if not open_file_with_default_app(path):
            AlertPopup('File Saved', f'Saved to:\n{path}').open()

    def _handle_export_select(self, export_type):
        self.export_dropdown.dismiss()
        if export_type == 'cloud':
            self._send_to_cloud()
            return
        if export_type == 'signout':
            sign_out()
            self._set_sync('Signed out', error=True)
            AlertPopup('Cloud', 'Signed out.\nYour notes stay on this device.').open()
            return
        runner = {
            'excel': (self._run_excel_export, 'Select Date Range for Excel'),
            'docx': (self._run_docx_export, 'Select Date Range for Word'),
            'pdf': (self._run_pdf_export, 'Select Date Range for PDF'),
        }[export_type]
        DateRangePopup(runner[1], self.store.communities, self.selected_community, runner[0]).open()

    @staticmethod
    def _range_text(start_d, end_d, community):
        return f'{start_d} to {end_d}' + (f' ({community})' if community else '')

    # --- cloud ---
    def _cloud_ready(self):
        """URL present? Otherwise tell the person once."""
        if load_cloud_settings()['url'].startswith(CLOUD_URL_PREFIX):
            return True
        AlertPopup('Cloud', 'This build has no web app URL.\nSet CLOUD_URL at the top of main.py.').open()
        return False

    # --- Google sign-in (the Google client comes from the script) ---
    def _start_sign_in(self, on_done):
        if self._signing_in or not self._cloud_ready():
            return
        self._signing_in = True
        threading.Thread(target=self._device_start_worker, args=(load_cloud_settings()['url'], on_done),
                         daemon=True).start()

    def _device_start_worker(self, url, on_done):
        info, s, cfg = None, None, None
        try:
            cfg, err = fetch_config(url)
            if not err and cfg['mode'] == 'google':
                s = load_cloud_settings()
                info, err = google_device_start(s['client_id'])
        except Exception as e:
            log_error('Sign-in start failed', sys.exc_info())
            info, err = None, f'{e}'
        Clock.schedule_once(lambda dt: self._device_started(s, on_done, info, err, cfg), 0)

    def _email_entered(self, email, on_done):
        self._signing_in = False
        if not save_cloud_settings(email=email):
            self._persist_failed_alert()
        on_done()

    def _ask_password(self, on_done):
        """The daily password: asked once a day, then remembered until the date changes."""
        if self._asking_password:
            return
        self._asking_password = True

        def entered(pw):
            self._asking_password = False
            if not save_cloud_settings(password=pw):
                self._persist_failed_alert()
            on_done()
        PromptPopup('Cloud password', 'Enter today\'s cloud password.',
                    'password', lambda t: t.strip() if t.strip() == today_password() else '',
                    entered, lambda: setattr(self, '_asking_password', False), secret=True).open()

    # --- first run: password + email on one screen ---
    def _first_run_check(self, dt):
        s = load_cloud_settings()
        if s['signed_in'] or self._asking_credentials or not self._cloud_ready():
            return
        self._asking_credentials = True
        CredentialsPopup(self._credentials_entered,
                         lambda: setattr(self, '_asking_credentials', False)).open()

    def _credentials_entered(self, password, email):
        if not save_cloud_settings(password=password):
            self._persist_failed_alert()
        url = load_cloud_settings()['url']
        threading.Thread(target=self._credentials_worker, args=(url, email), daemon=True).start()

    def _credentials_worker(self, url, email):
        cfg, err = None, None
        try:
            cfg, err = fetch_config(url)
        except Exception as e:
            log_error('First-run config failed', sys.exc_info())
            err = f'{e}'
        Clock.schedule_once(lambda dt: self._credentials_configured(cfg, err, email), 0)

    def _credentials_configured(self, cfg, err, email):
        self._asking_credentials = False
        if err:
            AlertPopup('Cloud', err if err != PASSWORD_NEEDED else 'Password was not accepted.').open()
            return
        if cfg['mode'] == 'email':
            # fetch_config may have reset the settings, so the address is saved after it
            if not save_cloud_settings(email=email):
                self._persist_failed_alert()
            self.refresh_access(silent=True)
        else:
            self._start_sign_in(lambda: self.refresh_access(silent=True))

    def _device_started(self, s, on_done, info, err, cfg=None):
        if err == PASSWORD_NEEDED:
            self._signing_in = False
            save_cloud_settings(password='')
            self._ask_password(lambda: self._start_sign_in(on_done))
            return
        if err:
            self._signing_in = False
            AlertPopup('Sign-in Error', err).open()
            return
        if cfg and cfg['mode'] == 'email':       # simple mode: no Google page, just the address
            PromptPopup('Your Gmail address',
                        'Enter the Gmail address that is listed\nfor you in the Access sheet.',
                        'name@gmail.com', clean_email, lambda email: self._email_entered(email, on_done),
                        lambda: setattr(self, '_signing_in', False)).open()
            return
        popup = SignInPopup(info['user_code'], info['verification_url'])
        popup.open()
        threading.Thread(target=self._device_poll_worker, args=(s, info, popup, on_done),
                         daemon=True).start()

    def _device_poll_worker(self, s, info, popup, on_done):
        deadline = time.time() + min(int(info.get('expires_in', 300)), 900)
        interval = max(int(info.get('interval', 5)), 1)
        result = None
        while time.time() < deadline and not popup.cancelled:
            time.sleep(interval)
            if popup.cancelled:
                break
            try:
                state, payload = google_device_poll(s['client_id'], s['client_secret'], info['device_code'])
            except Exception as e:
                state, payload = 'error', f'{e}'
            if state == 'pending':
                continue
            if state == 'slow':
                interval += 5
                continue
            result = (state, payload)
            break
        Clock.schedule_once(lambda dt: self._device_finished(popup, on_done, result), 0)

    def _device_finished(self, popup, on_done, result):
        self._signing_in = False
        cancelled = popup.cancelled
        popup.dismiss()
        if cancelled:
            return
        if result is None:
            AlertPopup('Sign-in', 'The code expired before you finished.\nTry again.').open()
            return
        state, payload = result
        if state != 'done':
            AlertPopup('Sign-in Error', payload).open()
            return
        email = str(id_token_info(payload['id_token']).get('email', ''))
        if not save_cloud_settings(refresh_token=payload['refresh_token'], email=email):
            self._persist_failed_alert()
        _remember_tokens(payload, payload['refresh_token'])
        on_done()

    # --- the Access list and background reading ---
    def refresh_access(self, silent=True):
        """Re-read the community list (and who may write where) from the Access sheet. Always quiet."""
        s = load_cloud_settings()
        if not (s['url'] and s['signed_in'] and s['password_ok']):
            return
        threading.Thread(target=self._access_worker, args=(s['url'],), daemon=True).start()

    def _access_worker(self, url):
        communities, err = None, None
        try:
            token, err = get_id_token()
            if not err:
                communities, _email, err = fetch_access(url, token)
        except Exception as e:
            log_error('Access list read failed', sys.exc_info())
            err = f'{e}'
        Clock.schedule_once(lambda dt: self._access_loaded(communities, err), 0)

    def _apply_access(self, communities):
        if self._editor_open:
            return                           # do not change communities under an open note
        self.store.set_communities([n for n, _w in communities], {n for n, w in communities if w})
        self._set_spinner(self.selected_community)
        self.create_calendar()
        self._update_status()

    def _access_loaded(self, communities, err):
        if err:
            self._set_sync('Signed out - use Sync' if err in (SIGN_IN_NEEDED, PASSWORD_NEEDED)
                           else 'Offline - showing the saved copy', error=True)
            return
        self._apply_access(communities)
        if self.selected_community:
            self._background_pull(self.selected_community)

    def _background_pull(self, community):
        s = load_cloud_settings()
        if not (s['url'] and s['signed_in'] and s['password_ok']):
            return
        if self._syncing:
            return                           # a Sync is running; it reads the sheet itself
        threading.Thread(target=self._bg_pull_worker, args=(s['url'], community, self._sync_gen),
                         daemon=True).start()

    def _bg_pull_worker(self, url, community, gen):
        notes, skipped, err = None, 0, None
        try:
            token, err = get_id_token()
            if not err:
                notes, skipped, err = pull_from_cloud(url, token, community)
        except Exception as e:
            log_error('Background read failed', sys.exc_info())
            err = f'{e}'
        Clock.schedule_once(lambda dt: self._bg_pulled(community, notes, err, gen), 0)

    def _bg_pulled(self, community, notes, err, gen):
        if self._syncing or gen != self._sync_gen:
            return                           # a Sync ran meanwhile: this read is out of date
        if err:
            self._set_sync('Offline - showing the saved copy' if err not in (SIGN_IN_NEEDED, PASSWORD_NEEDED)
                           else 'Signed out - use Sync', error=True)
            return
        if self._editor_open or community not in self.store.communities:
            return
        self.store.sync_from_sheet(community, notes)
        if community == self.selected_community:
            self.create_calendar()
            self._mark_synced()
            if self._sync_error:
                self._set_sync('')           # the problem is gone
            else:
                self._update_status()

    # --- Sync: runs quietly, then ONE message with an OK button ---
    def _send_to_cloud(self):
        if not self._cloud_ready():
            return
        s = load_cloud_settings()
        if not s['password_ok']:
            self._ask_password(self._send_to_cloud)
            return
        if not s['signed_in']:
            self._start_sign_in(self._send_to_cloud)
            return
        if self._syncing and time.time() - self._sync_started < 120:
            return                           # already running: ignore extra taps
        self._syncing = True
        self._sync_started = time.time()
        self._sync_gen += 1
        community = self.selected_community
        threading.Thread(target=self._cloud_pull_worker, args=(s['url'], community), daemon=True).start()

    def _sync_finished(self):
        self._syncing = False
        self._sync_gen += 1

    def _finish_sync(self, community, message):
        """Update the Sync button first; the message appears only when it has turned green."""
        self._sync_finished()
        self.create_calendar()
        if community and self.store.can_write(community) and self.store.has_pending(community):
            self._set_sync('Not fully synced - tap Sync again', error=True)
        else:
            self._mark_synced()
            self._set_sync(message)

    def _cloud_pull_worker(self, url, community):
        token, access, sheet_notes, skipped, err = None, None, None, 0, None
        try:
            token, err = get_id_token()
            if not err:
                access, _email, err = fetch_access(url, token)
            if not err:
                names = {n: w for n, w in access}
                if community not in names:
                    community = next(iter(names), None)
                if community is None:
                    err = 'No communities are listed for your Gmail in the Access sheet.'
                elif not names[community]:
                    err = (f'You can only view "{community}".\nYour Gmail is not listed for it in the Access sheet.')
                else:
                    sheet_notes, skipped, err = pull_from_cloud(url, token, community)
        except Exception as e:
            log_error('Cloud read failed', sys.exc_info())
            err = f'Read failed:\n{e}'
        Clock.schedule_once(
            lambda dt: self._cloud_pulled(url, token, community, access, sheet_notes, skipped, err), 0)

    def _cloud_pulled(self, url, token, community, access, sheet_notes, skipped, err):
        if access:
            self._apply_access(access)
        if err == PASSWORD_NEEDED:
            self._sync_finished()
            save_cloud_settings(password='')
            self._ask_password(self._send_to_cloud)
            return
        if err == SIGN_IN_NEEDED:
            self._sync_finished()
            self._start_sign_in(self._send_to_cloud)
            return
        if err:  # nothing is changed or uploaded unless the sheet could be read
            self._sync_finished()
            AlertPopup('Cloud', f'Nothing was changed or uploaded.\n{err}').open()
            return
        stats = self.store.sync_from_sheet(community, sheet_notes)
        self.create_calendar()
        if getattr(sheet_notes, 'tab_found', None) is False:     # worksheet missing: needs attention
            self._sync_finished()
            AlertPopup('Cloud', describe_sync(stats, skipped, sheet_notes, community)).open()
            return
        existing = len(sheet_notes)
        rows = self.store.pending_rows(community, sheet_notes)
        if not rows:
            self._finish_sync(community, self._sync_message(stats, 0, existing))
            return
        threading.Thread(target=self._cloud_push_worker,
                         args=(url, token, community, rows, stats, existing), daemon=True).start()

    @staticmethod
    def _sync_message(stats, uploaded, existing):
        """One short line shown under the community name after a sync."""
        def notes(n):
            return f'{n} note' + ('' if n == 1 else 's')
        parts = []
        if uploaded:
            parts.append(f'{notes(uploaded)} added to existing {notes(existing)}.')
        received = stats['added'] + stats['updated'] + stats['merged']
        if received:
            parts.append(f'{notes(received)} received from the sheet.')
        return ' '.join(parts) if parts else f'Up to date - {notes(existing)}.'

    def _cloud_push_worker(self, url, token, community, rows, stats, existing):
        verify = None
        try:
            ok, msg, saved = save_to_cloud(url, token, community, rows)
            if ok:      # read the sheet again: what it really holds decides what counts as synced
                verify, _skipped, _err = pull_from_cloud(url, token, community)
        except Exception as e:
            log_error('Cloud upload failed', sys.exc_info())
            ok, msg, saved = False, f'{e}', []
        Clock.schedule_once(
            lambda dt: self._cloud_done(ok, msg, community, rows, saved, stats, existing, verify), 0)

    def _cloud_done(self, ok, msg, community, rows, saved, stats, existing, verify=None):
        if not ok:
            self._sync_finished()
            self.create_calendar()
            AlertPopup('Cloud Error', f'Upload failed:\n{msg}').open()
            return
        done_dates = set(saved)
        if not self.store.mark_uploaded(community, [(d, t) for d, t in rows if d in done_dates]):
            self._persist_failed_alert()
        if verify is not None:
            self.store.sync_from_sheet(community, verify)    # marks whatever the sheet now holds as saved
        self._finish_sync(community, self._sync_message(stats, len(rows), existing))
        if len(saved) < len(rows):          # some notes were refused: say which
            AlertPopup('Cloud', msg).open()

    def _run_excel_export(self, start_d, end_d, community):
        rows = self.store.query(start_d, end_d, community)
        if not rows:
            AlertPopup('Export Excel',
                       f'No notes found for {self._range_text(start_d, end_d, community)}.').open()
            return
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font
            wb = Workbook()
            ws = wb.active
            ws.title = 'Annals Notes'
            single = community is not None
            if single:
                # Community name goes in a header row; no Community column needed
                comm_cell = ws.cell(row=1, column=1, value=f'Community: {community}')
                comm_cell.font = Font(bold=True)
                ws.append(['Date', 'Note'])
                for cell in ws[2]:
                    cell.font = Font(bold=True)
                for date_str, _comm, note in rows:
                    ws.append([date_str, note])
                ws.column_dimensions['A'].width = 14
                ws.column_dimensions['B'].width = 98
                data_start_row = 3
            else:
                ws.append(['Date', 'Community', 'Note'])
                for cell in ws[1]:
                    cell.font = Font(bold=True)
                for date_str, comm, note in rows:
                    ws.append([date_str, comm, note])
                ws.column_dimensions['A'].width = 14
                ws.column_dimensions['B'].width = 18
                ws.column_dimensions['C'].width = 80
                data_start_row = 2
            for row in ws.iter_rows(min_row=data_start_row):
                for cell in row:
                    cell.alignment = Alignment(wrap_text=True, vertical='top')

            if not _try_release_file(EXCEL_FILE):
                AlertPopup('File In Use',
                           'annals_notes.xlsx is open in Excel.\n'
                           'Please close it and export again.').open()
                return
            wb.save(EXCEL_FILE)
            self._deliver(EXCEL_FILE)
        except Exception as e:
            log_error('Excel Export Failed', sys.exc_info())
            AlertPopup('Export Error', f'Failed to save Excel file:\n{e}').open()

    def _run_docx_export(self, start_d, end_d, community):
        rows = self.store.query(start_d, end_d, community)
        if not rows:
            AlertPopup('Export Word',
                       f'No notes found for {self._range_text(start_d, end_d, community)}.').open()
            return
        try:
            if not _try_release_file(DOCX_FILE):
                AlertPopup('File In Use',
                           'annals_notes.docx is open in Word.\n'
                           'Please close it and export again.').open()
                return
            build_docx(DOCX_FILE, start_d, end_d, community, rows)
            self._deliver(DOCX_FILE)
        except Exception as e:
            log_error('Word Export Failed', sys.exc_info())
            AlertPopup('Export Error', f'Failed to save Word file:\n{e}').open()

    def _run_pdf_export(self, start_d, end_d, community):
        rows = self.store.query(start_d, end_d, community)
        if not rows:
            AlertPopup('Export PDF',
                       f'No notes found for {self._range_text(start_d, end_d, community)}.').open()
            return
        try:
            if not _try_release_file(PDF_FILE):
                AlertPopup('File In Use',
                           'annals_notes.pdf is open in a PDF viewer.\n'
                           'Please close it and export again.').open()
                return
            build_pdf(PDF_FILE, start_d.strftime('%d %b %Y'), end_d.strftime('%d %b %Y'), community, rows)
            self._deliver(PDF_FILE)
        except Exception as e:
            log_error('PDF Export Failed', sys.exc_info())
            AlertPopup('Export Error', f'Failed to save PDF file:\n{e}').open()


class MainApp(App):
    def build(self):
        self.title = 'Annals'
        return AnnalsNotesLayout()


class ErrorApp(App):
    def __init__(self, error_text, **kwargs):
        super().__init__(**kwargs)
        self.error_text = error_text

    def build(self):
        from kivy.uix.scrollview import ScrollView
        layout = BoxLayout(orientation='vertical', padding=15, spacing=10)
        layout.add_widget(Label(
            text="App Startup Error",
            font_size='22sp',
            bold=True,
            size_hint_y=0.1,
            color=(1, 0.3, 0.3, 1)
        ))
        sv = ScrollView(size_hint_y=0.9)
        lbl = Label(
            text=self.error_text,
            font_size='14sp',
            size_hint_y=None,
            halign='left',
            valign='top'
        )
        lbl.bind(texture_size=lambda instance, value: setattr(lbl, 'height', value[1]))
        lbl.bind(width=lambda instance, value: setattr(lbl, 'text_size', (value[0], None)))
        sv.add_widget(lbl)
        layout.add_widget(sv)
        return layout


if __name__ == '__main__':
    try:
        MainApp().run()
    except Exception:
        err = traceback.format_exc()
        try:
            log_error('Fatal startup error', sys.exc_info())
        except Exception:
            pass
        try:
            ErrorApp(err).run()
        except Exception:
            pass

